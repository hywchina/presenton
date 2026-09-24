import asyncio
import base64
import io
import os
import tempfile
from unittest.mock import AsyncMock

from docx import Document
from fastapi import HTTPException, UploadFile
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import pytest

from models.generated_file_plan import (
    GeneratedDocumentPlan,
    GeneratedDocumentSectionPlan,
    GeneratedFileType,
    GeneratedPresentationPlan,
    GeneratedPresentationSlidePlan,
)
from models.presentation_outline_model import SlideOutlineModel
from services.local_file_generation_service import (
    GeneratedReportMetadata,
    LocalFileGenerationService,
    PreparedInputImage,
)
from utils.outline_utils import get_images_for_slides_from_outline


def _png_bytes(color: str = "red") -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (32, 24), color=color).save(output, format="PNG")
    return output.getvalue()


def _prepared_image(tmp_path, index: int = 0) -> PreparedInputImage:
    path = tmp_path / f"image-{index}.png"
    data = _png_bytes()
    path.write_bytes(data)
    return PreparedInputImage(
        index=index,
        original_name=path.name,
        data=data,
        mime_type="image/png",
        path=str(path),
        asset_url=f"/app_data/images/{path.name}",
    )


def _document_plan() -> GeneratedDocumentPlan:
    return GeneratedDocumentPlan(
        title="离线分析报告",
        summary="本报告由本地模型生成。",
        sections=[
            GeneratedDocumentSectionPlan(
                heading="分析结果",
                paragraphs=["图片展示了一个红色方块。"],
                bullets=["完全离线", "保留图片"],
                image_indices=[0],
            )
        ],
        image_captions=["示例图片"],
    )


def test_prepare_images_normalizes_and_stores_png(tmp_path, monkeypatch):
    monkeypatch.setenv("APP_DATA_DIRECTORY", str(tmp_path / "app-data"))
    upload_file = tempfile.SpooledTemporaryFile()
    upload_file.write(_png_bytes())
    upload_file.seek(0)
    upload = UploadFile(filename="source.jpg", file=upload_file)
    service = LocalFileGenerationService()

    images = asyncio.run(service.prepare_images([upload]))

    assert len(images) == 1
    assert images[0].mime_type == "image/png"
    assert images[0].asset_url.startswith("/app_data/images/")
    assert os.path.isfile(images[0].path)
    with Image.open(images[0].path) as normalized:
        assert normalized.size == (32, 24)


def test_prepare_images_rejects_oversized_upload_while_streaming(monkeypatch):
    monkeypatch.setattr(
        "services.local_file_generation_service.MAX_IMAGE_BYTES", 3
    )
    upload_file = tempfile.SpooledTemporaryFile()
    upload_file.write(b"1234")
    upload_file.seek(0)
    upload = UploadFile(filename="large.png", file=upload_file)

    with pytest.raises(HTTPException) as exc_info:
        asyncio.run(LocalFileGenerationService().prepare_images([upload]))

    assert exc_info.value.status_code == 413


def test_render_markdown_embeds_image_in_single_file(tmp_path, monkeypatch):
    monkeypatch.setenv("APP_DATA_DIRECTORY", str(tmp_path / "app-data"))
    image = _prepared_image(tmp_path)
    service = LocalFileGenerationService()

    generated = service.render_markdown(
        _document_plan(),
        [image],
        "report",
        GeneratedReportMetadata(
            project_name="示范项目",
            report_type="客室设计方案报告",
            requested_by="项目设计师",
            generated_date="2026/09/24",
        ),
    )
    content = open(generated.path, encoding="utf-8").read()

    assert generated.extension == ".md"
    assert "# 离线分析报告" in content
    assert "## 报告信息" in content
    assert "## 章节导航" in content
    assert "## 1. 分析结果" in content
    assert "*图 1 示例图片*" in content
    expected = base64.b64encode(image.data).decode("ascii")
    assert f"data:image/png;base64,{expected}" in content


def test_render_word_contains_text_and_image(tmp_path, monkeypatch):
    monkeypatch.setenv("APP_DATA_DIRECTORY", str(tmp_path / "app-data"))
    image = _prepared_image(tmp_path)
    service = LocalFileGenerationService()

    generated = service.render_word(_document_plan(), [image], "report")
    document = Document(generated.path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert generated.extension == ".docx"
    assert "离线分析报告" in text
    assert "图片展示了一个红色方块" in text
    assert len(document.inline_shapes) == 1


def test_presentation_markdown_assigns_unplaced_images_and_local_urls(tmp_path):
    images = [_prepared_image(tmp_path, 0), _prepared_image(tmp_path, 1)]
    plan = GeneratedPresentationPlan(
        title="演示",
        slides=[
            GeneratedPresentationSlidePlan(
                title="首页",
                content_markdown="概览",
                image_indices=[],
            ),
            GeneratedPresentationSlidePlan(
                title="详情",
                content_markdown="分析",
                image_indices=[1],
            ),
        ],
        image_captions=["图一", "图二"],
    )

    markdown = LocalFileGenerationService().presentation_markdown(plan, images)
    extracted = get_images_for_slides_from_outline(
        [SlideOutlineModel(content=value) for value in markdown]
    )

    assert extracted == [
        ["/app_data/images/image-0.png"],
        ["/app_data/images/image-1.png"],
    ]


def test_render_presentation_uses_each_image_once_and_keeps_cover_clean(
    tmp_path, monkeypatch
):
    monkeypatch.setenv("APP_DATA_DIRECTORY", str(tmp_path / "app-data"))
    images = [_prepared_image(tmp_path, 0), _prepared_image(tmp_path, 1)]
    plan = GeneratedPresentationPlan(
        title="离线演示",
        slides=[
            GeneratedPresentationSlidePlan(
                title="首页", content_markdown="项目概览", image_indices=[0]
            ),
            GeneratedPresentationSlidePlan(
                title="设计背景", content_markdown="基于输入资料整理。"
            ),
            GeneratedPresentationSlidePlan(
                title="方案说明", content_markdown="图片对应方案。", image_indices=[1]
            ),
        ],
        image_captions=["图一", "图二"],
    )

    generated = LocalFileGenerationService().render_presentation(
        plan,
        images,
        "presentation",
        GeneratedReportMetadata(
            project_name="示范项目",
            report_type="客室设计方案报告",
            generated_date="2026/09/24",
        ),
    )
    presentation = Presentation(generated.path)

    assert generated.extension == ".pptx"
    assert len(presentation.slides) == 3
    assert not [
        shape
        for shape in presentation.slides[0].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert sum(
        1
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ) == 2


def test_generate_document_plan_sends_image_content(monkeypatch, tmp_path):
    image = _prepared_image(tmp_path)
    generated_payload = _document_plan().model_dump()
    generate = AsyncMock(return_value=generated_payload)
    monkeypatch.setattr(
        "services.local_file_generation_service.generate_structured_with_schema_retries",
        generate,
    )
    monkeypatch.setattr(
        "services.local_file_generation_service.get_client", lambda **_kwargs: object()
    )
    monkeypatch.setattr(
        "services.local_file_generation_service.get_llm_config", lambda: {}
    )
    monkeypatch.setattr(
        "services.local_file_generation_service.get_model", lambda: "qwen3-vl"
    )

    plan = asyncio.run(
        LocalFileGenerationService().generate_document_plan(
            text="生成报告",
            images=[image],
            language="Chinese",
            output_type=GeneratedFileType.WORD,
        )
    )

    assert plan.title == "离线分析报告"
    messages = generate.await_args.kwargs["messages"]
    user_content = messages[1].content
    assert any(getattr(part, "mime_type", None) == "image/png" for part in user_content)
    assert generate.await_args.args[1] == "qwen3-vl"
    assert generate.await_args.kwargs["max_tokens"] == 2048


def test_generate_presentation_plan_constrains_exact_slide_count(monkeypatch, tmp_path):
    image = _prepared_image(tmp_path)
    payload = GeneratedPresentationPlan(
        title="离线演示",
        slides=[
            GeneratedPresentationSlidePlan(title="一", content_markdown="内容一"),
            GeneratedPresentationSlidePlan(title="二", content_markdown="内容二"),
        ],
    ).model_dump()
    generate = AsyncMock(return_value=payload)
    monkeypatch.setattr(
        "services.local_file_generation_service.generate_structured_with_schema_retries",
        generate,
    )
    monkeypatch.setattr(
        "services.local_file_generation_service.get_client", lambda **_kwargs: object()
    )
    monkeypatch.setattr(
        "services.local_file_generation_service.get_llm_config", lambda: {}
    )
    monkeypatch.setattr(
        "services.local_file_generation_service.get_model", lambda: "qwen3-vl"
    )

    plan = asyncio.run(
        LocalFileGenerationService().generate_presentation_plan(
            text="生成两页演示",
            images=[image],
            language="Chinese",
            n_slides=2,
        )
    )

    schema = generate.await_args.kwargs["json_schema"]
    assert len(plan.slides) == 2
    assert schema["properties"]["slides"]["minItems"] == 2
    assert schema["properties"]["slides"]["maxItems"] == 2
